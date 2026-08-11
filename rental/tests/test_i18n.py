"""Guards for the translation pipeline.

The catalogues are built and compiled by `manage.py trans_extract` /
`trans_compile` rather than gettext, so these tests check the things a missing
`msgfmt` would otherwise let through: incomplete locales, placeholders that
would raise at render time, and captions that reach openpyxl as lazy proxies.
"""

import gettext
import io
from pathlib import Path

from django.conf import settings
from django.test import TestCase
from django.test.utils import override_settings
from django.urls import reverse
from django.utils import translation
from openpyxl import load_workbook

from rental.i18n_tools import check_placeholders, parse_po
from rental.models import Booking
from rental.tests.factories import make_booking, make_car, make_company

LANGUAGES = [code for code, _label in settings.LANGUAGES]


def catalogue(language):
    return parse_po(
        Path(settings.LOCALE_PATHS[0]) / language / "LC_MESSAGES" / "django.po"
    )


class RestoreLanguageMixin:
    """LocaleMiddleware activates a language per request and never puts it back.

    Without this, a test that fetches /de/ leaves German active for every test
    that runs afterwards.
    """

    def setUp(self):
        super().setUp()
        self.addCleanup(translation.activate, settings.LANGUAGE_CODE)


class CatalogueTests(TestCase):
    def test_every_configured_language_has_a_catalogue(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                self.assertTrue(catalogue(language).messages)

    def test_every_language_is_fully_translated(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                messages = catalogue(language).messages
                missing = [m.msgid for m in messages if not m.translated]
                self.assertEqual(missing, [], f"{language} has untranslated strings")

    def test_no_translation_is_left_fuzzy(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                fuzzy = [m.msgid for m in catalogue(language).messages if m.fuzzy]
                self.assertEqual(fuzzy, [])

    def test_placeholders_survive_translation(self):
        # A msgstr naming a placeholder its msgid does not have raises KeyError
        # the moment the string is interpolated.
        for language in LANGUAGES:
            with self.subTest(language=language):
                problems = check_placeholders(catalogue(language))
                self.assertEqual(
                    [(src, tgt) for src, tgt, _missing, _extra in problems], []
                )

    def test_plurals_carry_both_forms(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                for message in catalogue(language).messages:
                    if message.is_plural:
                        self.assertEqual(sorted(message.msgstr_plural), [0, 1])

    def test_all_catalogues_cover_the_same_strings(self):
        reference = {m.key for m in catalogue(settings.LANGUAGE_CODE).messages}
        for language in LANGUAGES:
            with self.subTest(language=language):
                self.assertEqual({m.key for m in catalogue(language).messages}, reference)

    def test_compiled_mo_matches_the_po(self):
        # Compared by content, not mtime: a fresh checkout has arbitrary
        # timestamps but the two files must still agree.
        for language in LANGUAGES:
            with self.subTest(language=language):
                po = Path(settings.LOCALE_PATHS[0]) / language / "LC_MESSAGES" / "django.po"
                mo = po.with_suffix(".mo")
                self.assertTrue(mo.exists(), f"{language}: run manage.py trans_compile")

                with mo.open("rb") as fh:
                    compiled = gettext.GNUTranslations(fh)

                stale = [
                    m.msgid
                    for m in parse_po(po).messages
                    if not m.is_plural and compiled.gettext(m.msgid) != m.msgstr
                ]
                self.assertEqual(
                    stale, [], f"{language}: .mo is stale — run manage.py trans_compile"
                )


class ActivationTests(RestoreLanguageMixin, TestCase):
    def test_strings_differ_between_languages(self):
        rendered = {}
        for language in LANGUAGES:
            with translation.override(language):
                rendered[language] = translation.gettext("Find a car")

        self.assertEqual(rendered["en"], "Find a car")
        for language in ["sq", "it", "de"]:
            self.assertNotEqual(
                rendered[language], "Find a car", f"{language} fell back to English"
            )

    def test_plural_selection_works_per_language(self):
        for language in LANGUAGES:
            with self.subTest(language=language), translation.override(language):
                one = translation.ngettext("%(n)s seat", "%(n)s seats", 1)
                many = translation.ngettext("%(n)s seat", "%(n)s seats", 5)
                self.assertNotEqual(one, many)
                self.assertIn("%(n)s", one)
                self.assertIn("%(n)s", many)


@override_settings(ALLOWED_HOSTS=["testserver"])
class LocalisedPageTests(RestoreLanguageMixin, TestCase):
    def test_home_page_renders_in_every_language(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                response = self.client.get(f"/{language}/")
                self.assertEqual(response.status_code, 200)
                body = response.content.decode("utf-8")

                with translation.override(language):
                    expected = translation.gettext("All rights reserved.")
                self.assertIn(expected, body)

    def test_language_prefix_sets_the_html_lang_attribute(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                response = self.client.get(f"/{language}/")
                self.assertContains(response, f'<html lang="{language}"')


class LanguageSwitcherTests(RestoreLanguageMixin, TestCase):

    def test_the_sprite_carries_a_symbol_for_every_language(self):
        body = self.client.get("/en/").content.decode("utf-8")
        for language in LANGUAGES:
            with self.subTest(language=language):
                self.assertIn(f'id="f-{language}"', body)

    def test_the_menu_lists_every_language_with_its_flag(self):
        body = self.client.get("/en/").content.decode("utf-8")
        for language in LANGUAGES:
            with self.subTest(language=language):
                self.assertIn(f'href="#f-{language}"', body)

    def test_the_pill_shows_the_flag_of_the_active_language(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                body = self.client.get(f"/{language}/").content.decode("utf-8")
                pill = body.split('class="lang"', 1)[1].split("</button>", 1)[0]
                self.assertIn(f'href="#f-{language}"', pill)

    def test_only_the_active_language_is_marked_current(self):
        body = self.client.get("/it/").content.decode("utf-8")
        rows = body.split('class="lang-option"')[1:]
        current = [r for r in rows if 'aria-current="true"' in r.split(">", 1)[0]]
        self.assertEqual(len(current), 1)
        self.assertIn("#f-it", current[0])

    def test_a_language_without_artwork_falls_back_to_the_globe(self):
        from rental.templatetags.ui import flag

        self.assertIn("#f-en", flag("en"))
        self.assertIn("#i-globe", flag("fr"))
        self.assertIn("#i-globe", flag(""))

    def test_regional_codes_resolve_to_their_base_flag(self):
        from rental.templatetags.ui import flag

        self.assertIn("#f-en", flag("en-gb"))


class LocalisedExportTests(RestoreLanguageMixin, TestCase):

    def setUp(self):
        super().setUp()
        self.company = make_company("Delta Rentals")
        make_booking(make_car(self.company), status=Booking.Status.CONFIRMED)
        self.client.force_login(self.company.owner)

    def _export(self, key):
        return self.client.get(reverse("company_report_export", args=[key]))

    def test_every_report_exports_in_every_language(self):
        for language in LANGUAGES:
            for key in ["bookings", "revenue", "fleet", "customers", "maintenance"]:
                with self.subTest(language=language, report=key):
                    with translation.override(language):
                        response = self._export(key)
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.content[:2], b"PK")

    def test_sheet_names_and_headers_follow_the_active_language(self):
        with translation.override("de"):
            response = self._export("bookings")
            expected_sheet = translation.gettext("Bookings")
            expected_header = translation.gettext("Reference")

        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn(expected_sheet, workbook.sheetnames)

        sheet = workbook[expected_sheet]
        headers = [cell.value for cell in sheet[5]]
        self.assertIn(expected_header, headers)

    def test_headers_are_written_as_text_not_proxies(self):
        with translation.override("it"):
            response = self._export("fleet")
            sheet_name = translation.gettext("Fleet")

        sheet = load_workbook(io.BytesIO(response.content))[sheet_name]
        for cell in sheet[5]:
            if cell.value is not None:
                self.assertIsInstance(cell.value, str)
