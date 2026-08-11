import io
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rental.models import AuditLog, Booking
from rental.tests.factories import make_booking, make_car, make_company, make_superuser


class CompanyReportTests(TestCase):
    def setUp(self):
        self.company = make_company("Delta Rentals")
        self.car = make_car(self.company, price="45.00")
        self.booking = make_booking(self.car, status=Booking.Status.CONFIRMED)
        self.client.force_login(self.company.owner)

    def _workbook(self, response):
        return load_workbook(io.BytesIO(response.content))

    def test_bookings_export_is_a_real_xlsx(self):
        response = self.client.get(reverse("company_report_export", args=["bookings"]))

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(response.content[:2], b"PK")

        workbook = self._workbook(response)
        self.assertIn("Bookings", workbook.sheetnames)

    def test_export_contains_the_booking_data(self):
        response = self.client.get(reverse("company_report_export", args=["bookings"]))
        sheet = self._workbook(response)["Bookings"]
        values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
        self.assertIn(self.booking.reference, values)

    def test_revenue_report_has_multiple_sheets(self):
        response = self.client.get(reverse("company_report_export", args=["revenue"]))
        workbook = self._workbook(response)
        for name in ["Summary", "Daily", "By vehicle", "Bookings"]:
            self.assertIn(name, workbook.sheetnames)

    def test_every_company_report_downloads(self):
        for key in ["bookings", "revenue", "fleet", "customers", "maintenance"]:
            with self.subTest(report=key):
                response = self.client.get(
                    reverse("company_report_export", args=[key])
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.content[:2], b"PK")

    def test_unknown_report_is_a_404(self):
        response = self.client.get(reverse("company_report_export", args=["nonsense"]))
        self.assertEqual(response.status_code, 404)

    def test_exports_are_audited(self):
        self.client.get(reverse("company_report_export", args=["bookings"]))
        self.assertTrue(
            AuditLog.objects.filter(
                action=AuditLog.Action.REPORT_EXPORTED, company=self.company
            ).exists()
        )

    def test_date_filter_narrows_the_export(self):
        response = self.client.get(
            reverse("company_report_export", args=["bookings"]),
            {"from": "1999-01-01", "to": "1999-12-31"},
        )
        sheet = self._workbook(response)["Bookings"]
        values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
        self.assertNotIn(self.booking.reference, values)


class PlatformReportTests(TestCase):
    def setUp(self):
        self.company_a = make_company("Alpha Rentals")
        self.company_b = make_company("Beta Rentals")
        make_booking(make_car(self.company_a), status=Booking.Status.CONFIRMED)
        make_booking(make_car(self.company_b), status=Booking.Status.CONFIRMED)

        self.root = make_superuser()
        self.client.force_login(self.root)

    def test_every_platform_report_downloads(self):
        for key in [
            "companies", "bookings", "revenue", "fleet",
            "customers", "users", "activity",
        ]:
            with self.subTest(report=key):
                response = self.client.get(
                    reverse("superadmin_report_export", args=[key])
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.content[:2], b"PK")

    def test_platform_export_spans_all_companies(self):
        response = self.client.get(
            reverse("superadmin_report_export", args=["companies"])
        )
        sheet = load_workbook(io.BytesIO(response.content))["Companies"]
        values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
        self.assertIn("Alpha Rentals", values)
        self.assertIn("Beta Rentals", values)

    def test_company_filter_narrows_a_platform_export(self):
        response = self.client.get(
            reverse("superadmin_report_export", args=["bookings"]),
            {"company": self.company_a.pk},
        )
        sheet = load_workbook(io.BytesIO(response.content))["Bookings"]
        values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]

        a_ref = Booking.objects.get(company=self.company_a).reference
        b_ref = Booking.objects.get(company=self.company_b).reference
        self.assertIn(a_ref, values)
        self.assertNotIn(b_ref, values)

    def test_empty_report_still_produces_a_valid_file(self):
        Booking.objects.all().delete()
        response = self.client.get(
            reverse("superadmin_report_export", args=["bookings"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content[:2], b"PK")
